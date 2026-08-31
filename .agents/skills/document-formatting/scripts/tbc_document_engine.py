import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

class TBCDocumentEngine:
    """
    مُحرّك أتمتة الخطابات والمراسلات لشركة TBC ومؤسسة إعمار الفرعة.
    يعتمد على قواعد التنسيق المحفوظة في بطاقات الذاكرة المرجعية.
    """

    @staticmethod
    def _set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    @staticmethod
    def _set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    @staticmethod
    def _set_font(run, font_name="Sakkal Majalla", size_pt=14, bold=False, color=None):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.bold = bold
        if color:
            run.font.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)
        rFonts.set(qn('w:cs'), font_name)
        rPr.append(rFonts)
        
        rtl = OxmlElement('w:rtl')
        rtl.set(qn('w:val'), '1')
        rPr.append(rtl)

    @classmethod
    def _add_para(cls, container, text="", bold=False, size=14, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=4, color=None):
        p = container.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement('w:bidi')
        bidi.set(qn('w:val'), '1')
        pPr.append(bidi)
        if text:
            run = p.add_run(text)
            cls._set_font(run, font_name="Sakkal Majalla", size_pt=size, bold=bold, color=color)
        return p

    @staticmethod
    def _make_table_rtl(table):
        tblPr = table._tbl.tblPr
        bidiVisual = OxmlElement('w:bidiVisual')
        tblPr.append(bidiVisual)

    @classmethod
    def _fill_cell_rtl(cls, cell, text, bold=False, size=14, align=None):
        p = cell.paragraphs[0]
        if align is not None:
            p.alignment = align
        p._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
        run = p.add_run(text)
        cls._set_font(run, size_pt=size, bold=bold)

    @classmethod
    def generate_guarantee_letter(
        cls,
        company_name="شركة / ..........................................................................",
        factory_type="(مصنع / شركة صناعة الكابلات المعتمدة)",
        addressee=": شركة تطوير للمباني TBC / شركة سماء للاستشارات الهندسية",
        product=": كابلات الجهد المنخفض",
        quantity=": كامل المبنى",
        project_name="مشروع الصيانة الطارئة لمدرسة أم المؤمنين زينب الأسدية للطفولة المبكرة والروضة الملحقة",
        po_number="TBC006687",
        so_number="305333",
        contractor="مؤسسة إعمار الفرعة للمقاولات العامة",
        guarantee_years="10",
        output_path="خطاب_ضمان.docx",
        single_page=True
    ):
        doc = Document()
        
        # Margins setup (0.8 inches for single page compact mode, 0.98 for standard)
        margin_size = 0.8 if single_page else 0.98
        cell_margin_v = 40 if single_page else 100
        
        for section in doc.sections:
            section.top_margin = Inches(margin_size)
            section.bottom_margin = Inches(margin_size)
            section.left_margin = Inches(margin_size)
            section.right_margin = Inches(margin_size)
            
        # RTL base
        styles = doc.styles
        normal_style = styles['Normal']
        pPr = normal_style._element.get_or_add_pPr()
        bidi = OxmlElement('w:bidi')
        bidi.set(qn('w:val'), '1')
        pPr.append(bidi)

        # 1. Header
        header_font_size = 16 if single_page else 18
        cls._add_para(doc, company_name, bold=True, size=header_font_size, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, color=RGBColor(31, 73, 125))
        cls._add_para(doc, factory_type, bold=False, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6 if single_page else 12)
        cls._add_para(doc, "وزارة التعليم - شركة تطوير للمباني TBC", bold=True, size=header_font_size, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10 if single_page else 14)

        # 2. Addressee Info Table
        info_table = doc.add_table(rows=3, cols=2)
        cls._make_table_rtl(info_table)
        
        rows_info = [
            ("إلى", addressee),
            ("المنتج", product),
            ("الكمية", quantity)
        ]
        for i, (k, v) in enumerate(rows_info):
            r = info_table.rows[i]
            c0 = r.cells[0]
            c0.width = Inches(1.0)
            cls._fill_cell_rtl(c0, k, bold=True, size=14)
            c1 = r.cells[1]
            c1.width = Inches(4.5)
            cls._fill_cell_rtl(c1, v, bold=True, size=14)

        cls._add_para(doc, "", space_after=6 if single_page else 8)

        # 3. Project Details Section
        cls._add_para(doc, "(بيانات المشروع)", bold=True, size=15 if single_page else 16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6 if single_page else 8)

        table_data = [
            ("اسم المشروع", project_name),
            ("رقم طلب الشراء", po_number),
            ("رقم طلب البيع", so_number),
            ("المقاول المنفذ", contractor),
            ("مدة الضمان", f"({guarantee_years}) سنوات"),
            ("بداية فترة الضمان", "من تاريخ الاستلام الابتدائي للمشروع"),
            ("نهاية فترة الضمان", f"({guarantee_years}) سنوات من تاريخ الاستلام الابتدائي للمشروع")
        ]

        table = doc.add_table(rows=len(table_data), cols=2)
        table.autofit = False
        cls._make_table_rtl(table)

        tbl_font_size = 13 if single_page else 14
        for i, (label, val) in enumerate(table_data):
            row = table.rows[i]
            
            cell_lbl = row.cells[0]
            cell_lbl.width = Inches(2.2)
            cls._set_cell_background(cell_lbl, "FFFFFF")
            cls._set_cell_margins(cell_lbl, top=cell_margin_v, bottom=cell_margin_v, left=100, right=100)
            cls._fill_cell_rtl(cell_lbl, label, bold=True, size=tbl_font_size, align=WD_ALIGN_PARAGRAPH.RIGHT)
            
            cell_val = row.cells[1]
            cell_val.width = Inches(4.5)
            cls._set_cell_background(cell_val, "FFFFFF")
            cls._set_cell_margins(cell_val, top=cell_margin_v, bottom=cell_margin_v, left=100, right=100)
            cls._fill_cell_rtl(cell_val, val, bold=True, size=tbl_font_size, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Table Borders
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            '  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            '  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            '  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            '  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            '  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            '  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            '</w:tblBorders>'
        )
        table._tbl.tblPr.append(borders)

        cls._add_para(doc, "", space_after=8 if single_page else 12)

        # 4. Guarantee Declaration Body
        body_text = (
            "نحن شركة .......................................................................... نتعهد بضمان "
            "ومطابقة المنتج المذكور أعلاه حسب شروط ومواصفات وزارة التعليم للمشروع المذكور "
            f"والعينة المعتمدة من الجهة المختصة بالوزارة أعلاه لمدة {guarantee_years} سنوات وذلك ضد عيوب التصنيع والتجميع، وكذلك "
            "نتعهد بإصلاح أو تبديل المنتج في حال وجود أي خلل فني فيه في مدة أقصاها ثلاث أيام من تاريخ إبلاغنا بذلك "
            "من الإدارة المذكورة أعلاه، ولا يشمل هذا الضمان سوء الاستخدام أو أي صيانة من شركة أخرى غير الفريق الفني لشركتنا."
        )
        cls._add_para(doc, body_text, bold=False, size=14 if single_page else 15, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=14 if single_page else 24)

        # 5. Signatures Table
        sig_table = doc.add_table(rows=1, cols=2)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cls._make_table_rtl(sig_table)
        
        row_sig = sig_table.rows[0]
        c_factory = row_sig.cells[0]
        c_factory.width = Inches(3.36)
        c_contractor = row_sig.cells[1]
        c_contractor.width = Inches(3.44)

        def append_sig_para(cell, idx, text, bold=True):
            if idx < len(cell.paragraphs):
                p = cell.paragraphs[idx]
            else:
                p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(8 if single_page else 12)
            p._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
            run = p.add_run(text)
            cls._set_font(run, size_pt=14, bold=bold)

        # Factory Signatures
        append_sig_para(c_factory, 0, "المصنع / الشركة المنتجة:", bold=True)
        append_sig_para(c_factory, 1, "شركة ..............................................................", bold=False)
        append_sig_para(c_factory, 2, "التوقيع: ...........................................", bold=False)
        append_sig_para(c_factory, 3, "التاريخ:        /        / 1448 هـ", bold=False)
        append_sig_para(c_factory, 4, "الختم:", bold=False)

        # Contractor Signatures
        append_sig_para(c_contractor, 0, "المقاول المنفذ:", bold=True)
        append_sig_para(c_contractor, 1, contractor, bold=True)
        append_sig_para(c_contractor, 2, "التوقيع: ...........................................", bold=False)
        append_sig_para(c_contractor, 3, "التاريخ:        /        / 1448 هـ", bold=False)
        append_sig_para(c_contractor, 4, "الختم:", bold=False)

        doc.save(output_path)
        return output_path
