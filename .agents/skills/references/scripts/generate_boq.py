# -*- coding: utf-8 -*-
"""
سكريبت توليد جداول الكميات (BOQ) وتحليل الأسعار بصيغة Excel
تأليف: وكيل boq-analyst
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def generate_boq_excel(project_name, items_data, output_path):
    """
    توليد ملف جدول الكميات والتسعير
    items_data: قائمة بقواميس البنود:
    [
        {"num": 1, "desc": "تكسير وإزالة السور القديم...", "unit": "م.ط", "qty": 500, "unit_price": 50},
        ...
    ]
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "جدول الكميات"
    ws.views.sheetView[0].showGridLines = True
    
    # ضبط الاتجاه RTL
    ws.sheet_view.rightToLeft = True
    
    font_family = "Traditional Arabic"
    title_font = Font(name=font_family, size=16, bold=True, color="1F497D")
    header_font = Font(name=font_family, size=12, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=11)
    bold_data_font = Font(name=font_family, size=11, bold=True)
    
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    stripe_fill = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
    total_fill = PatternFill(start_color="E6EDF5", end_color="E6EDF5", fill_type="solid")
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)
    
    thin_border_side = Side(border_style="thin", color="D3D3D3")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(bottom=Side(border_style="medium", color="1F497D"))
    
    # 1. ترويسة العنوان
    ws.merge_cells("A1:G1")
    ws["A1"] = f"جدول الكميات والتسعير لمشروع: {project_name}"
    ws["A1"].font = title_font
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 40
    
    # 2. عناوين الجدول
    headers = ["رقم البند", "بيان الأعمال (وصف البند)", "الوحدة", "الكمية", "سعر الوحدة (ريال)", "السعر الإجمالي (ريال)", "الكمية كتابة"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border
        
    ws.row_dimensions[3].height = 28
    
    # 3. ملء بنود جدول الكميات
    row_idx = 4
    for item in items_data:
        qty = item.get("qty", 0)
        unit_price = item.get("unit_price", 0)
        
        ws.cell(row=row_idx, column=1, value=item.get("num", row_idx - 3)).alignment = align_center
        ws.cell(row=row_idx, column=2, value=item.get("desc", "")).alignment = align_right
        ws.cell(row=row_idx, column=3, value=item.get("unit", "عدد")).alignment = align_center
        ws.cell(row=row_idx, column=4, value=qty).alignment = align_center
        ws.cell(row=row_idx, column=5, value=unit_price).alignment = align_center
        
        # استخدام صيغة Excel لحساب السعر الإجمالي
        total_formula = f"=D{row_idx}*E{row_idx}"
        ws.cell(row=row_idx, column=6, value=total_formula).alignment = align_center
        
        ws.cell(row=row_idx, column=7, value=item.get("qty_words", "")).alignment = align_center
        
        is_stripe = (row_idx % 2 == 0)
        for col_idx in range(1, 8):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if is_stripe:
                cell.fill = stripe_fill
                
            # تنسيق أرقام العملات والكميات
            if col_idx in [4, 5]:
                cell.number_format = '#,##0.00'
            elif col_idx == 6:
                cell.number_format = '#,##0.00'
                
        ws.row_dimensions[row_idx].height = 24
        row_idx += 1
        
    # 4. خلاصة الحسابات (المجموع، الضريبة، الإجمالي النهائي)
    ws.cell(row=row_idx, column=2, value="المجموع الإجمالي الخاضع للضريبة").alignment = align_right
    ws.cell(row=row_idx, column=6, value=f"=SUM(F4:F{row_idx-1})").alignment = align_center
    
    ws.cell(row=row_idx+1, column=2, value="ضريبة القيمة المضافة (15%)").alignment = align_right
    ws.cell(row=row_idx+1, column=6, value=f"=F{row_idx}*0.15").alignment = align_center
    
    ws.cell(row=row_idx+2, column=2, value="الإجمالي النهائي شامل ضريبة القيمة المضافة").alignment = align_right
    ws.cell(row=row_idx+2, column=6, value=f"=F{row_idx}+F{row_idx+1}").alignment = align_center
    
    # تطبيق التنسيق على صفوف المجاميع
    for r in range(row_idx, row_idx + 3):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.cell(row=r, column=2).font = bold_data_font
        ws.cell(row=r, column=6).font = bold_data_font
        ws.cell(row=r, column=6).number_format = '#,##0.00'
        
        for c in range(1, 8):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border
            cell.fill = total_fill
            
        ws.row_dimensions[r].height = 24
        
    # 5. ملاءمة عرض الأعمدة
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 3:
                continue
            if cell.value:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["G"].width = 20
    
    # حفظ الملف
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path

if __name__ == "__main__":
    test_items = [
        {"num": 1, "desc": "أعمال تكسير وإزالة السور القديم وترحيل المخلفات للمرمى العمومي.", "unit": "م.ط", "qty": 500, "unit_price": 45, "qty_words": "خمسمائة"},
        {"num": 2, "desc": "أعمال حفر لزوم القواعد والأساسات في تربة صخرية أو عادية لعمق 1.5م.", "unit": "م3", "qty": 588, "unit_price": 20, "qty_words": "خمسمائة وثمان وثمانون"},
        {"num": 3, "desc": "توريد وصب خرسانة عادية أسفل الأساسات إجهاد 250 كجم/سم2.", "unit": "م3", "qty": 24, "unit_price": 280, "qty_words": "أربعة وعشرون"},
        {"num": 4, "desc": "توريد وصب خرسانة مسلحة مقاومة للكبريتات إجهاد 350 كجم/سم2 لزوم الأساسات والميد ورقاب الأعمدة والطبان.", "unit": "م3", "qty": 60, "unit_price": 420, "qty_words": "ستون"}
    ]
    generate_boq_excel("إعادة تأهيل سور مقبرة الرفيعة", test_items, "g:/M.Yonis/.agents/skills/references/scripts/test_boq.xlsx")
    print("تم توليد جدول كميات التجربة بنجاح في: test_boq.xlsx")
