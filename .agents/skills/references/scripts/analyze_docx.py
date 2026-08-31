# -*- coding: utf-8 -*-
"""
سكريبت تحليل ملفات DOCX واستخراج الهيكل والتنسيق تلقائياً
تأليف: وكيل reference-calibrator
"""

import sys
import os
from docx import Document

def analyze_docx(file_path, output_report_path):
    if not os.path.exists(file_path):
        print(f"Error: File not found at {file_path}")
        return False
        
    try:
        doc = Document(file_path)
        report = []
        report.append(f"# تقرير تحليل مستند مرجعي: {os.path.basename(file_path)}")
        report.append(f"- **المسار الكامل:** {file_path}")
        report.append(f"- **عدد الفقرات:** {len(doc.paragraphs)}")
        report.append(f"- **عدد الجداول:** {len(doc.tables)}")
        report.append("\n---\n")
        
        # 1. تحليل العناوين والأقسام الرئيسية
        report.append("## 1. الهيكل العام والعناوين المستخرجة\n")
        headings_found = []
        for idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if not text:
                continue
                
            # التحقق من أنماط العناوين
            style_name = p.style.name
            is_heading = False
            level = 0
            
            if style_name.startswith('Heading') or style_name.startswith('Title') or style_name.startswith('Subtitle'):
                is_heading = True
                level = style_name[-1] if style_name[-1].isdigit() else 1
            elif len(text) < 100 and (text.startswith(tuple(f"{i}." for i in range(1, 30))) or text.startswith(tuple(f"{i}-" for i in range(1, 30)))):
                is_heading = True
                level = 1
                
            if is_heading:
                indent = "  " * (int(level) - 1) if str(level).isdigit() else ""
                report.append(f"{indent}- **[العنوان {style_name}]:** {text}")
                headings_found.append((text, style_name))
                
        if not headings_found:
            report.append("لم يتم العثور على عناوين واضحة بالأنماط القياسية. قد يكون النص منسقاً يدوياً.")
            
        report.append("\n---\n")
        
        # 2. تحليل عينات من الفقرات والتنسيقات
        report.append("## 2. تحليل التنسيقات والخطوط\n")
        # فحص عينات من الخطوط المستخدمة
        fonts = set()
        font_sizes = set()
        for p in doc.paragraphs[:50]: # فحص أول 50 فقرة كعينة
            for run in p.runs:
                if run.font.name:
                    fonts.add(run.font.name)
                if run.font.size:
                    font_sizes.add(run.font.size.pt)
                    
        report.append(f"- **الخطوط المكتشفة في العينة:** {', '.join(fonts) if fonts else 'غير محددة (الوضع الافتراضي)'}")
        report.append(f"- **أحجام الخطوط المكتشفة:** {', '.join(str(s) for s in font_sizes) if font_sizes else 'غير محددة'}")
        
        # فحص الهوامش
        if doc.sections:
            sec = doc.sections[0]
            report.append("- **هوامش الصفحة المكتشفة (بوصة):**")
            report.append(f"  - الهامش العلوي: {sec.top_margin.inches:.2f}")
            report.append(f"  - الهامش السفلي: {sec.bottom_margin.inches:.2f}")
            report.append(f"  - الهامش الأيمن: {sec.right_margin.inches:.2f}")
            report.append(f"  - الهامش الأيسر: {sec.left_margin.inches:.2f}")
            
        report.append("\n---\n")
        
        # 3. تحليل الجداول وهيكلها
        report.append("## 3. تحليل الجداول المضمنة\n")
        for t_idx, table in enumerate(doc.tables, 1):
            rows_count = len(table.rows)
            cols_count = len(table.columns) if rows_count > 0 else 0
            report.append(f"### الجدول رقم {t_idx} (أبعاد: {rows_count} صف × {cols_count} عمود)")
            
            # جلب ترويسة الجدول (الصف الأول)
            if rows_count > 0:
                headers = [cell.text.strip().replace('\n', ' ') for cell in table.rows[0].cells]
                # إزالة التكرار الناتج عن خلايا مدمجة أفقياً
                clean_headers = []
                for h in headers:
                    if not clean_headers or clean_headers[-1] != h:
                        clean_headers.append(h)
                report.append(f"- **عناوين الأعمدة:** { ' | '.join(clean_headers) }")
            
            # جلب عينة من البيانات (أول صفين بعد الترويسة)
            if rows_count > 1:
                report.append("- **عينة من البيانات:**")
                for r in range(1, min(3, rows_count)):
                    row_data = [cell.text.strip().replace('\n', ' ') for cell in table.rows[r].cells]
                    clean_row = []
                    for rd in row_data:
                        if not clean_row or clean_row[-1] != rd:
                            clean_row.append(rd)
                    report.append(f"  - الصف {r}: { ' / '.join(clean_row[:5]) } ...")
            report.append("")
            
        # حفظ التقرير
        os.makedirs(os.path.dirname(output_report_path), exist_ok=True)
        with open(output_report_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(report))
            
        print(f"Success: Report generated at {output_report_path}")
        return True
    except Exception as e:
        print(f"Error during analysis: {str(e)}")
        return False

if __name__ == "__main__":
    if len(sys.argv) < 3:
        # اختبار افتراضي
        analyze_docx(
            r"G:\M.Yonis\خرج\منهجية انجاز الأعمال لمشروع تشغيل وصيانة مباني ومرافق بلدية الخرج.docx",
            r"g:\M.Yonis\.agents\skills\references\scripts\analysis_report.txt"
        )
    else:
        analyze_docx(sys.argv[1], sys.argv[2])
