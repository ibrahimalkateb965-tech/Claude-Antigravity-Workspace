#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Prompt Library Curator & Sync Engine
Author: Antigravity Multi-Agent Core
Description: Programmatically formats, categorizes, and appends engineered prompts 
into 'مكتبة الأوامر.xlsx' with precise cell styling, fonts, borders, and auto-row detection.
"""

import os
import sys
import json
import argparse
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

DEFAULT_EXCEL_PATH = r"G:\أرشيف_ونسخ_احتياطية\ملفات_مكررة_وعامة\AI\مكتبة الأوامر.xlsx"

# Standard Cell Styles matching 'مكتبة الأوامر.xlsx'
THIN_SIDE = Side(border_style="thin", color="000000")
CELL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

FONT_NO = Font(name="Segoe UI", size=10, bold=False, color="000000")
FONT_NAME = Font(name="Segoe UI", size=10, bold=True, color="000000")
FONT_PROMPT = Font(name="Consolas", size=10, bold=True, color="000000")
FONT_DESC = Font(name="Segoe UI", size=10, bold=False, color="000000")

ALIGN_NO = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_NAME = Alignment(horizontal="right", vertical="center", wrap_text=True)
ALIGN_PROMPT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_DESC = Alignment(horizontal="right", vertical="center", wrap_text=True)

# Sheet keywords for intelligent routing
SHEET_ROUTING = {
    "CLI_Core.Claude": ["cli", "claude", "antigravity", "fleet", "agcli", "clear", "compact", "context", "orchestrator", "commander", "advisor"],
    "Workflows_Pipelines.Claude": ["workflow", "pipeline", "ci", "cd", "testflight", "release", "github actions", "build", "automation"],
    "Gates_Hooks.Claude": ["hook", "gate", "quality gate", "خطاف", "بوابة", "audit", "مراجعة", "فحص"],
    "Custom_Skills.Claude": ["skill", "مهارة", "agent", "وكيل", "subagent", "role", "تخصيص"],
    "Config_Environments.Claude": ["config", "environment", "settings", "بيئة", "إعدادات", "gradle", "sdk", "path", "memory", "ram"],
    "Construction_MCP_Suite": ["boq", "مخطط", "عقد", "مقاولات", "هندسة", "كميات", "tbc", "مستخلص", "docx", "إنشاءات"]
}

def detect_best_sheet(name, prompt, desc, available_sheets):
    text = f"{name} {prompt} {desc}".lower()
    scores = {}
    for sheet, keywords in SHEET_ROUTING.items():
        if sheet in available_sheets:
            score = sum(1 for kw in keywords if kw in text)
            scores[sheet] = score
            
    best_sheet = max(scores, key=scores.get)
    if scores[best_sheet] > 0:
        return best_sheet
    return "CLI_Core.Claude" if "CLI_Core.Claude" in available_sheets else "prompts"

def add_prompt_to_library(name, prompt, desc="", sheet_name=None, excel_path=None):
    excel_path = excel_path or DEFAULT_EXCEL_PATH
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found at: {excel_path}")

    # Open workbook
    wb = openpyxl.load_workbook(excel_path)
    
    if not sheet_name or sheet_name not in wb.sheetnames:
        sheet_name = detect_best_sheet(name, prompt, desc, wb.sheetnames)
        
    ws = wb[sheet_name]

    # Find the next clean row
    # Scan from bottom to find last populated row in Col B or Col C
    target_row = 1
    max_no = 0
    for r in range(ws.max_row, 0, -1):
        v_b = ws.cell(r, 2).value
        v_c = ws.cell(r, 3).value
        v_no = ws.cell(r, 1).value
        if v_b or v_c:
            target_row = r + 1
            break

    # Determine next 'No'
    for r in range(1, target_row):
        val = ws.cell(r, 1).value
        if isinstance(val, int):
            if val > max_no:
                max_no = val
        elif isinstance(val, str) and val.strip().isdigit():
            num = int(val.strip())
            if num > max_no:
                max_no = num

    next_no = max_no + 1

    # Insert cells
    c_no = ws.cell(row=target_row, column=1, value=next_no)
    c_name = ws.cell(row=target_row, column=2, value=name)
    c_prompt = ws.cell(row=target_row, column=3, value=prompt)
    c_desc = ws.cell(row=target_row, column=4, value=desc if desc else None)

    # Apply precise styling
    c_no.font = FONT_NO
    c_no.alignment = ALIGN_NO
    c_no.border = CELL_BORDER

    c_name.font = FONT_NAME
    c_name.alignment = ALIGN_NAME
    c_name.border = CELL_BORDER

    c_prompt.font = FONT_PROMPT
    c_prompt.alignment = ALIGN_PROMPT
    c_prompt.border = CELL_BORDER

    c_desc.font = FONT_DESC
    c_desc.alignment = ALIGN_DESC
    c_desc.border = CELL_BORDER

    # Set row height if needed
    ws.row_dimensions[target_row].height = 45

    # Attempt save with fallback
    try:
        wb.save(excel_path)
    except PermissionError:
        # File is locked in Excel - try COM automation
        try:
            import win32com.client
            xl = win32com.client.GetObject(Class="Excel.Application")
            target_wb = None
            for w in xl.Workbooks:
                if "مكتبة الأوامر" in w.Name:
                    target_wb = w
                    break
            if target_wb:
                target_sheet = target_wb.Sheets(sheet_name)
                target_sheet.Cells(target_row, 1).Value = next_no
                target_sheet.Cells(target_row, 2).Value = name
                target_sheet.Cells(target_row, 3).Value = prompt
                target_sheet.Cells(target_row, 4).Value = desc if desc else ""
                target_wb.Save()
            else:
                raise PermissionError("File locked by Excel and workbook not found in running instances.")
        except Exception as e:
            raise PermissionError(f"Cannot save Excel file because it is open in Microsoft Excel. Please save/close it and retry. Error: {e}")

    return {
        "status": "success",
        "excel_path": excel_path,
        "sheet_name": sheet_name,
        "row_number": target_row,
        "no": next_no,
        "name": name,
        "prompt": prompt,
        "desc": desc
    }

def main():
    parser = argparse.ArgumentParser(description="Prompt Library Curator & Sync Engine")
    parser.add_argument("--name", "-n", default=None, help="Command / Prompt Name (Arabic + English)")
    parser.add_argument("--prompt", "-p", default=None, help="The prompt text or command trigger")
    parser.add_argument("--desc", "-d", default="", help="Description & Usage Notes")
    parser.add_argument("--sheet", "-s", default=None, help="Target Sheet / Tab name")
    parser.add_argument("--file", "-f", default=DEFAULT_EXCEL_PATH, help="Path to Excel workbook")
    parser.add_argument("--list-sheets", action="store_true", help="List all available sheets in workbook")
    parser.add_argument("--json", action="store_true", help="Output result as JSON")
    args = parser.parse_args()

    try:
        excel_path = args.file or DEFAULT_EXCEL_PATH
        if args.list_sheets:
            if not os.path.exists(excel_path):
                raise FileNotFoundError(f"Excel file not found at: {excel_path}")
            wb = openpyxl.load_workbook(excel_path, read_only=True)
            sheets = wb.sheetnames
            if args.json:
                print(json.dumps({"sheets": sheets}, ensure_ascii=False, indent=2))
            else:
                print(f"📑 Available Sheets in '{os.path.basename(excel_path)}' ({len(sheets)} sheets):")
                for i, s in enumerate(sheets, 1):
                    print(f"  {i}. {s}")
            return

        if not args.name or not args.prompt:
            parser.error("the following arguments are required: --name/-n, --prompt/-p (unless using --list-sheets)")

        res = add_prompt_to_library(
            name=args.name,
            prompt=args.prompt,
            desc=args.desc,
            sheet_name=args.sheet,
            excel_path=args.file
        )
        if args.json:
            print(json.dumps(res, ensure_ascii=False, indent=2))
        else:
            print(f"✅ Successfully added to Prompts Library!")
            print(f"📑 Sheet: {res['sheet_name']}")
            print(f"📍 Location: Row #{res['row_number']} (No: {res['no']})")
            print(f"🏷️ Name: {res['name']}")
    except Exception as e:
        if args.json:
            print(json.dumps({"status": "error", "error": str(e)}, ensure_ascii=False))
        else:
            print(f"❌ Error: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
